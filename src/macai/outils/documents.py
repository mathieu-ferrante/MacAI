"""Documents : lire tout format, décrire une image, créer Word / Excel / PowerPoint / CSV."""
from __future__ import annotations

import base64
import csv
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from .. import extraction
from .base import ErreurOutil, Resultat, Risque, outil
from .fichiers import resoudre, sans_collision


def destination(chemin: str, extension: str) -> Path:
    """Nom seul → sur le Bureau ; extension ajoutée si absente ; jamais d'écrasement."""
    if "/" not in chemin and "\\" not in chemin and not chemin.startswith("~"):
        chemin = f"~/Desktop/{chemin}"
    p = resoudre(chemin, ecriture=True)
    if p.suffix.lower() != extension:
        p = p.with_name(p.name + extension)
    p.parent.mkdir(parents=True, exist_ok=True)
    return sans_collision(p)


@outil("lire_document", "Lit le contenu texte d'un fichier : PDF (même scanné), Word, Excel, PowerPoint, "
       "texte, code, et images (texte visible + description du contenu).", "fichiers",
       {"chemin": ("string", "chemin du fichier"),
        "debut": ("integer", "caractère de départ pour lire la suite d'un long document, facultatif")},
       requis=["chemin"], mac_seulement=False)
def lire_document(ctx, chemin, debut=0):
    p = resoudre(chemin, doit_exister=True)
    if p.is_dir():
        raise ErreurOutil(f"{p} est un dossier : utilise lister_dossier.")
    if p.suffix.lower() not in extraction.EXTRACTIBLES:
        raise ErreurOutil(f"Format {p.suffix or 'inconnu'} non lisible. Essaie ouvrir_fichier.")
    texte = extraction.extraire(p)
    if not texte.strip():
        return Resultat(True, f"{p.name} : aucun texte trouvé.")
    debut = max(0, int(debut or 0))
    morceau = texte[debut:debut + 12000]
    suite = f"\n[… suite disponible : relire avec debut={debut + 12000}, {len(texte)} caractères au total]" \
        if len(texte) > debut + 12000 else ""
    return Resultat(True, f"{p.name} :\n{morceau}{suite}", {"fichier": str(p), "longueur": len(texte)})


def _image_reduite(p: Path) -> bytes:
    """Réduit l'image à 1024 px (moins de calcul pour le modèle de vision)."""
    if sys.platform == "darwin":
        with tempfile.TemporaryDirectory() as d:
            sortie = Path(d) / "img.jpg"
            r = subprocess.run(["sips", "-Z", "1024", "-s", "format", "jpeg", str(p), "--out", str(sortie)],
                               capture_output=True, timeout=30)
            if r.returncode == 0 and sortie.exists():
                return sortie.read_bytes()
    return p.read_bytes()


def demander_vision(ctx, image: bytes, question: str) -> str:
    """Interroge le modèle de vision. S'il est distinct de l'agent, l'agent est déchargé d'abord (8 Go de RAM) ;
    avec un modèle multimodal unique (défaut), rien à décharger ni à recharger."""
    cfg = ctx.config.modeles
    if cfg["vision"] != cfg["agent"]:
        ctx.ollama.decharger(cfg["agent"])
    rep = ctx.ollama.discuter(cfg["vision"], [{"role": "user", "content": question,
                                               "images": [base64.b64encode(image).decode()]}])
    return rep.texte.strip()


@outil("decrire_image", "Décrit précisément une image ou répond à une question sur elle (modèle de vision local).",
       "fichiers", {"chemin": ("string", "chemin de l'image"),
                    "question": ("string", "question précise, facultatif")}, requis=["chemin"], mac_seulement=False)
def decrire_image(ctx, chemin, question=""):
    p = resoudre(chemin, doit_exister=True)
    if p.suffix.lower() not in extraction.IMAGES:
        raise ErreurOutil(f"{p.name} n'est pas une image.")
    reponse = demander_vision(ctx, _image_reduite(p), question or "Décris cette image en français, en détail.")
    return f"{p.name} : {reponse}"


# --- Création -------------------------------------------------------------------

def _gras(paragraphe, texte: str) -> None:
    for i, morceau in enumerate(re.split(r"\*\*(.+?)\*\*", texte)):
        if morceau:
            paragraphe.add_run(morceau).bold = i % 2 == 1


@outil("creer_document_word", "Crée un document Word (.docx). Le contenu s'écrit en Markdown simple : "
       "# Titre, ## Sous-titre, - puce, 1. liste, **gras**, tableaux | a | b |.", "fichiers",
       {"chemin": ("string", "nom ou chemin du fichier, ex. ~/Desktop/Résumé.docx"),
        "contenu": ("string", "contenu en Markdown simple")}, risque=Risque.CREATION, mac_seulement=False,
       annonce=lambda a: f"Créer le document Word {a['chemin']}")
def creer_document_word(ctx, chemin, contenu):
    import docx
    p = destination(chemin, ".docx")
    d = docx.Document()
    tableau: list[list[str]] = []

    def vider_tableau():
        if not tableau:
            return
        t = d.add_table(rows=len(tableau), cols=max(len(l) for l in tableau))
        t.style = "Table Grid"
        for i, ligne in enumerate(tableau):
            for j, cellule in enumerate(ligne):
                t.cell(i, j).text = cellule
                if i == 0:
                    for run in t.cell(i, j).paragraphs[0].runs:
                        run.bold = True
        tableau.clear()

    for ligne in contenu.splitlines():
        s = ligne.strip()
        if s.startswith("|") and s.endswith("|"):
            cellules = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cellules):
                tableau.append(cellules)
            continue
        vider_tableau()
        if not s:
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)", s):
            d.add_heading(m.group(2), level=len(m.group(1)) - 1)  # # → titre, ## → niveau 1…
        elif m := re.match(r"^[-*•]\s+(.*)", s):
            _gras(d.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\d+[.)]\s+(.*)", s):
            _gras(d.add_paragraph(style="List Number"), m.group(1))
        else:
            _gras(d.add_paragraph(), s)
    vider_tableau()
    d.save(str(p))
    return Resultat(True, f"Document Word créé : {p}", {"fichier": str(p)})


NOMBRE = re.compile(r"^-?[\d\s  ]+([.,]\d+)?\s*(€|%|\$)?$")
# Gardés en texte : zéro initial (code postal, numéro de téléphone, matricule) — « 06000 », « 06 12 34 56 78 ».
ZERO_INITIAL = re.compile(r"^-?0\d")
TELEPHONE = re.compile(r"^\+?\d{1,3}([\s.\-]\d{2}){3,}$")
FORMULE = ("=", "+", "-", "@")


def _valeur(v):
    """« 1 284,50 € » → 1284.5 ; « 06000 » et « 06 12 34 56 78 » restent du texte ; le reste inchangé."""
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v).strip()
    if ZERO_INITIAL.match(s) or TELEPHONE.match(s):
        return s
    if NOMBRE.match(s) and any(c.isdigit() for c in s):
        propre = re.sub(r"[\s  €$%]", "", s).replace(",", ".")
        try:
            n = float(propre)
            return int(n) if n.is_integer() and "." not in propre else n
        except ValueError:
            pass
    return s


def _sans_formule(v):
    """Injection de formules (CSV ouvert dans Excel) : « =HYPERLINK(…) », « +cmd… » → précédés d'une apostrophe."""
    if isinstance(v, str) and v[:1] in FORMULE and not NOMBRE.match(v.strip()):
        return "'" + v
    return v


def _decouper_ligne(texte: str) -> list[str]:
    """« Marie;30 », « Marie\t30 » ou « Marie, 30 » → ["Marie", "30"]."""
    for sep in (";", "\t", ","):
        if sep in texte:
            return [c.strip() for c in texte.split(sep)]
    return [texte.strip()]


def _tableau(colonnes, lignes) -> tuple[list[str], list[list]]:
    """Accepte ce qu'un petit modèle produit : lignes en listes, en dictionnaires {colonne: valeur} ou en
    chaînes « Marie;30 » ; colonnes en liste ou en chaîne."""
    if isinstance(colonnes, str):
        colonnes = _decouper_ligne(colonnes)
    colonnes = [str(c) for c in (colonnes or [])]
    if isinstance(lignes, (str, dict)):
        lignes = lignes.splitlines() if isinstance(lignes, str) else [lignes]
    sortie = []
    for ligne in lignes or []:
        if isinstance(ligne, dict):
            if not colonnes:
                colonnes = [str(k) for k in ligne]
            minuscules = {str(k).strip().lower(): v for k, v in ligne.items()}
            sortie.append([ligne.get(c, minuscules.get(c.strip().lower(), "")) for c in colonnes])
        elif isinstance(ligne, str):
            sortie += [_decouper_ligne(l) for l in ligne.splitlines() if l.strip()]
        elif isinstance(ligne, (list, tuple)):
            sortie.append(list(ligne))
        else:
            sortie.append([ligne])
    return colonnes, sortie


@outil("creer_tableur", "Crée un fichier Excel (.xlsx) à partir de colonnes et de lignes. Les nombres "
       "(même « 1 284,50 € ») sont convertis pour pouvoir être calculés.", "fichiers",
       {"chemin": ("string", "nom ou chemin, ex. ~/Desktop/Factures.xlsx"),
        "colonnes": {"type": "array", "items": {"type": "string"}, "description": "titres des colonnes"},
        "lignes": {"type": "array", "items": {"type": "array", "items": {}}, "description": "lignes de valeurs"},
        "total": ("boolean", "ajouter une ligne de total pour les colonnes numériques")},
       requis=["chemin", "colonnes", "lignes"], risque=Risque.CREATION, mac_seulement=False,
       annonce=lambda a: f"Créer le tableur {a['chemin']}")
def creer_tableur(ctx, chemin, colonnes, lignes, total=False):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    colonnes, lignes = _tableau(colonnes, lignes)
    p = destination(chemin, ".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Données"
    ws.append(list(colonnes))
    for c in ws[1]:
        c.data_type = "s"  # un titre « =… » reste du texte
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
    for ligne in lignes:
        ws.append([_valeur(v) for v in ligne])
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value.startswith(FORMULE):
                c.data_type = "s"  # texte, jamais une formule exécutée par Excel
    n = len(lignes)
    if total and n:
        ws.append([])
        ligne_total = ["Total"] + [""] * (len(colonnes) - 1)
        for j in range(1, len(colonnes)):
            if all(isinstance(ws.cell(i, j + 1).value, (int, float)) for i in range(2, n + 2)
                   if ws.cell(i, j + 1).value is not None):
                lettre = get_column_letter(j + 1)
                ligne_total[j] = f"=SUM({lettre}2:{lettre}{n + 1})"
        ws.append(ligne_total)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for j, titre in enumerate(colonnes, 1):
        largeur = max([len(str(titre))] + [len(str(ws.cell(i, j).value or "")) for i in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(j)].width = min(largeur + 2, 60)
    ws.freeze_panes = "A2"
    wb.save(str(p))
    return Resultat(True, f"Tableur créé : {p} ({n} lignes)", {"fichier": str(p)})


@outil("creer_csv", "Crée un fichier CSV (séparateur ; pour Excel en français).", "fichiers",
       {"chemin": ("string", "nom ou chemin"),
        "colonnes": {"type": "array", "items": {"type": "string"}, "description": "titres"},
        "lignes": {"type": "array", "items": {"type": "array", "items": {}}, "description": "lignes"}},
       risque=Risque.CREATION, mac_seulement=False)
def creer_csv(ctx, chemin, colonnes, lignes):
    colonnes, lignes = _tableau(colonnes, lignes)
    p = destination(chemin, ".csv")
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([_sans_formule(c) for c in colonnes])
        w.writerows([[_sans_formule(v) for v in ligne] for ligne in lignes])
    return Resultat(True, f"CSV créé : {p}", {"fichier": str(p)})


@outil("creer_presentation", "Crée une présentation PowerPoint (.pptx) : une diapositive de titre puis une "
       "diapositive par élément (titre + points).", "fichiers",
       {"chemin": ("string", "nom ou chemin"), "titre": ("string", "titre de la présentation"),
        "diapositives": {"type": "array", "description": "liste de diapositives",
                         "items": {"type": "object", "properties": {
                             "titre": {"type": "string"},
                             "points": {"type": "array", "items": {"type": "string"}}}}}},
       risque=Risque.CREATION, mac_seulement=False, annonce=lambda a: f"Créer la présentation {a['chemin']}")
def creer_presentation(ctx, chemin, titre, diapositives):
    from pptx import Presentation
    from pptx.util import Pt
    p = destination(chemin, ".pptx")
    prs = Presentation()
    d0 = prs.slides.add_slide(prs.slide_layouts[0])
    d0.shapes.title.text = titre
    for diapo in diapositives:
        if isinstance(diapo, str):
            diapo = {"titre": diapo, "points": []}
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = str(diapo.get("titre", ""))
        corps = s.placeholders[1].text_frame
        points = diapo.get("points") or []
        if isinstance(points, str):  # « - un\n- deux » ou une seule phrase
            points = [l.strip().lstrip("-*•").strip() for l in points.splitlines() if l.strip()]
        for i, point in enumerate(points):
            para = corps.paragraphs[0] if i == 0 else corps.add_paragraph()
            para.text = str(point)
            para.font.size = Pt(20)
    prs.save(str(p))
    return Resultat(True, f"Présentation créée : {p} ({len(diapositives) + 1} diapositives)", {"fichier": str(p)})
