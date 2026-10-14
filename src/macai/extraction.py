"""Extraction du texte de tous les formats courants (partagée par l'agent et l'indexation)."""
from __future__ import annotations

import csv
import io
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

TEXTE = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".xml", ".html",
         ".htm", ".css", ".js", ".ts", ".tsx", ".jsx", ".py", ".swift", ".java", ".kt", ".c", ".h", ".cpp",
         ".rs", ".go", ".rb", ".php", ".sh", ".zsh", ".sql", ".tex", ".log", ".ini", ".cfg", ".srt", ".vtt"}
IMAGES = {".jpg", ".jpeg", ".png", ".heic", ".webp", ".tiff", ".tif", ".bmp", ".gif"}
BUREAUTIQUE = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".rtf", ".doc", ".odt", ".pages", ".numbers", ".key"}
EXTRACTIBLES = TEXTE | IMAGES | BUREAUTIQUE


def extraire(chemin: Path, limite: int = 200_000, ocr: bool = True) -> str:
    """Texte d'un fichier ; chaîne vide si le format n'est pas pris en charge."""
    ext = chemin.suffix.lower()
    if ext in TEXTE:
        return _texte(chemin, limite)
    fonction = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
                ".rtf": _textutil, ".doc": _textutil, ".odt": _textutil,
                ".pages": _iwork, ".numbers": _iwork, ".key": _iwork}.get(ext)
    if fonction:
        texte = fonction(chemin, ocr) if fonction in (_pdf,) else fonction(chemin)
        return texte[:limite]
    if ext in IMAGES and ocr:
        return _image(chemin)[:limite]
    return ""


def _texte(chemin: Path, limite: int) -> str:
    brut = chemin.open("rb").read(limite * 2)
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return brut.decode(enc)[:limite]
        except UnicodeDecodeError:
            continue
    return ""


def _pdf(chemin: Path, ocr: bool = True) -> str:
    from pypdf import PdfReader
    try:
        lecteur = PdfReader(str(chemin))
        if lecteur.is_encrypted:
            try:
                lecteur.decrypt("")
            except Exception:  # noqa: BLE001
                return "[PDF protégé par mot de passe]"
        pages = []
        for i, page in enumerate(lecteur.pages[:300]):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(f"[page {i + 1}]\n{t}")
        texte = "\n\n".join(pages)
    except Exception as e:  # noqa: BLE001 - PDF corrompus fréquents
        texte = f"[PDF illisible : {e}]"
    # Peu de texte : c'est probablement un scan → OCR Apple Vision.
    if ocr and len(texte.strip()) < 100 and sys.platform == "darwin":
        from .mac.vision import ocr_pdf
        try:
            texte = ocr_pdf(chemin)
        except Exception:  # noqa: BLE001
            pass
    return texte


def _docx(chemin: Path) -> str:
    import docx
    d = docx.Document(str(chemin))
    parties = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for ligne in table.rows:
            parties.append(" | ".join(c.text.strip() for c in ligne.cells))
    return "\n".join(parties)


def _xlsx(chemin: Path, lignes_max: int = 300) -> str:
    import openpyxl
    classeur = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    sortie = []
    for feuille in classeur.worksheets:
        sortie.append(f"[feuille {feuille.title}]")
        tampon = io.StringIO()
        w = csv.writer(tampon, delimiter=";")
        for i, ligne in enumerate(feuille.iter_rows(values_only=True)):
            if i >= lignes_max:
                sortie.append(f"[… feuille tronquée après {lignes_max} lignes]")
                break
            if any(v is not None for v in ligne):
                w.writerow(["" if v is None else v for v in ligne])
        sortie.append(tampon.getvalue())
    classeur.close()
    return "\n".join(sortie)


def _pptx(chemin: Path) -> str:
    from pptx import Presentation
    sortie = []
    for i, diapo in enumerate(Presentation(str(chemin)).slides, 1):
        textes = [f.text_frame.text for f in diapo.shapes if f.has_text_frame and f.text_frame.text.strip()]
        sortie.append(f"[diapositive {i}]\n" + "\n".join(textes))
    return "\n\n".join(sortie)


def _textutil(chemin: Path) -> str:
    if sys.platform != "darwin":
        return ""
    r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(chemin)], capture_output=True,
                       text=True, timeout=60)
    return r.stdout


def _iwork(chemin: Path) -> str:
    """Pages/Numbers/Keynote : on lit l'aperçu PDF intégré au paquet s'il existe."""
    try:
        if chemin.is_dir():
            apercu = chemin / "QuickLook" / "Preview.pdf"
            return _pdf(apercu) if apercu.exists() else ""
        with zipfile.ZipFile(chemin) as z:
            for nom in ("QuickLook/Preview.pdf", "preview.pdf"):
                if nom in z.namelist():
                    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
                        f.write(z.read(nom))
                        tmp = Path(f.name)
                    try:
                        return _pdf(tmp)
                    finally:
                        tmp.unlink(missing_ok=True)
    except (zipfile.BadZipFile, OSError):
        pass
    return ""


def _image(chemin: Path) -> str:
    if sys.platform != "darwin":
        return ""
    from .mac import vision
    try:
        texte = vision.lire_texte(chemin)
        tags = vision.etiquettes(chemin)
    except Exception:  # noqa: BLE001
        return ""
    parties = []
    if tags:
        parties.append("Contenu de l'image : " + ", ".join(tags))
    if texte.strip():
        parties.append("Texte dans l'image :\n" + texte)
    return "\n".join(parties)
