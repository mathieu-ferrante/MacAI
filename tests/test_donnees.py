"""Données et garde-fous : documents, dates, index, MCP, serveur, journal, Trousseau, mémoire, compétences."""
import csv
import json
import os
import socket
import stat
import sys
import threading
import urllib.request
from pathlib import Path

import pytest

from macai.outils.base import ErreurOutil, Registre, Resultat, Risque, outil

SERVEUR_MCP = [sys.executable, str(Path(__file__).with_name("faux_serveur_mcp.py"))]


# --- Documents ------------------------------------------------------------------------------

def test_tableau_lignes_en_dict_chaines_et_listes():
    from macai.outils.documents import _tableau
    assert _tableau(["Nom", "Âge"], [{"Nom": "Marie", "Âge": 30}, {"nom": "Paul"}]) == (
        ["Nom", "Âge"], [["Marie", 30], ["Paul", ""]])
    assert _tableau([], [{"a": 1, "b": 2}]) == (["a", "b"], [[1, 2]])
    assert _tableau("Nom;Âge", ["Marie;30", "Paul\t40", "Zoé, 50"])[1] == [["Marie", "30"], ["Paul", "40"], ["Zoé", "50"]]
    assert _tableau(["Nom", "Âge"], "Marie;30\nPaul;40")[1] == [["Marie", "30"], ["Paul", "40"]]
    assert _tableau(["Nom"], ["Marie;30\nPaul;40"])[1] == [["Marie", "30"], ["Paul", "40"]]
    assert _tableau(["N"], [["a", 1], ("b", 2), 3])[1] == [["a", 1], ["b", 2], [3]]


def test_valeurs_texte_gardees():
    from macai.outils.documents import _valeur
    assert _valeur("06 12 34 56 78") == "06 12 34 56 78"
    assert _valeur("06000") == "06000"
    assert _valeur("01.23.45.67.89") == "01.23.45.67.89"
    assert _valeur("1 284,50 €") == 1284.5 and _valeur("0") == 0 and _valeur("0,5") == 0.5
    assert _valeur("-12") == -12


def test_csv_sans_injection_de_formule(ctx, tmp_path):
    from macai.outils.documents import creer_csv
    r = creer_csv(ctx, str(tmp_path / "a.csv"), ["Nom", "=Titre"],
                  [{"Nom": "=HYPERLINK(\"x\")", "=Titre": "-12"}, "@SUM(A1);+33 6 12"])
    with open(r.donnees["fichier"], encoding="utf-8-sig") as f:
        lignes = list(csv.reader(f, delimiter=";"))
    assert lignes == [["Nom", "'=Titre"], ["'=HYPERLINK(\"x\")", "-12"], ["'@SUM(A1)", "'+33 6 12"]]


def test_tableur_lignes_dict_et_formule_en_texte(ctx, tmp_path):
    import openpyxl
    from macai.outils.documents import creer_tableur
    r = creer_tableur(ctx, str(tmp_path / "t.xlsx"), ["Nom", "Tél", "Montant"],
                      [{"Nom": "=1+1", "Tél": "06 12 34 56 78", "Montant": "12,5"}, "Paul;06000;3"], total=True)
    ws = openpyxl.load_workbook(r.donnees["fichier"]).active
    assert [c.value for c in ws[2]] == ["=1+1", "06 12 34 56 78", 12.5]
    assert ws["A2"].data_type == "s"                       # texte, pas une formule
    assert ws["B3"].value == "06000" and ws["C3"].value == 3
    assert ws.cell(ws.max_row, 3).value == "=SUM(C2:C3)"   # le total, lui, reste une formule


def test_presentation_points_en_chaine(ctx, tmp_path):
    from pptx import Presentation
    from macai.outils.documents import creer_presentation
    r = creer_presentation(ctx, str(tmp_path / "p.pptx"), "Titre",
                           [{"titre": "A", "points": "- un\n- deux"}, {"titre": "B", "points": "seul"}])
    prs = Presentation(r.donnees["fichier"])
    textes = [[p.text for p in s.placeholders[1].text_frame.paragraphs] for s in list(prs.slides)[1:]]
    assert textes == [["un", "deux"], ["seul"]]


# --- Dates ------------------------------------------------------------------------------------

def test_lire_date_formes_parlees():
    from datetime import datetime, timedelta
    from macai.outils.organisation import lire_date
    demain = (datetime.now() + timedelta(days=1)).date()
    assert lire_date("demain à 14h") == datetime.combine(demain, datetime.min.time()).replace(hour=14)
    assert lire_date("demain matin").hour == 9 and lire_date("demain soir").hour == 19
    assert lire_date("demain après-midi").hour == 14 and lire_date("demain midi").hour == 12
    assert lire_date("2026-10-15T09:30:00Z") == datetime(2026, 10, 15, 9, 30)
    assert lire_date("2026-10-15 à 14h30") == datetime(2026, 10, 15, 14, 30)
    assert lire_date("15/10/2026 14h") == datetime(2026, 10, 15, 14, 0)
    with pytest.raises(ErreurOutil, match="AAAA-MM-JJ"):
        lire_date("la semaine des quatre jeudis")
    with pytest.raises(ErreurOutil):
        lire_date("demain 25h")


def test_date_illisible_message_clair(ctx):
    from macai.outils import organisation  # noqa: F401 - inscrit les outils
    from macai.outils.base import REGISTRE
    r = REGISTRE.get("creer_evenement").executer(ctx, {"titre": "x", "debut": "un jour"})
    assert not r.ok and "Erreur inattendue" not in r.texte and "Date non comprise" in r.texte
    assert REGISTRE.get("envoyer_mail").toujours_confirmer
    assert REGISTRE.get("preparer_mail").risque == Risque.CREATION


# --- Index -------------------------------------------------------------------------------------

def test_matrice_rechargee_apres_ecriture_d_un_autre_processus(tmp_path):
    from macai.recherche.index import Index
    from .conftest import FauxOllama
    lecteur = Index(tmp_path / "i.db", FauxOllama())
    ecrivain = Index(tmp_path / "i.db", FauxOllama())       # autre connexion, comme l'indexeur
    ecrivain.enregistrer(tmp_path / "a.txt", 1.0, 1, "document", "abeilles et miel")
    assert len(lecteur._vecteurs()[0]) == 1
    ecrivain.enregistrer(tmp_path / "b.txt", 1.0, 1, "document", "zèbres de savane")
    assert len(lecteur._vecteurs()[0]) == 2


def test_indexeur_n_oublie_que_les_fichiers_disparus(cfg, tmp_path, monkeypatch):
    from macai.recherche import indexeur
    from macai.recherche.index import Index
    present = tmp_path / "present.txt"
    present.write_text("x")
    index = Index(tmp_path / "i.db")
    index.enregistrer(present, 1.0, 1, "document", "x")
    index.enregistrer(tmp_path / "parti.txt", 1.0, 1, "document", "y")
    monkeypatch.setattr(indexeur, "parcourir", lambda *a, **k: iter(()))   # dossier illisible : rien vu
    _, disparus = indexeur.Indexeur(cfg, index).a_faire()
    assert disparus == [str(tmp_path / "parti.txt")]


def test_indexeur_ignore_les_fichiers_icloud_non_telecharges(tmp_path, monkeypatch):
    from macai.recherche import indexeur
    (tmp_path / "local.txt").write_text("x")
    (tmp_path / "nuage.txt").write_text("y")
    vrai_scandir = os.scandir

    class Entree:
        def __init__(self, e):
            self._e, self.name, self.path = e, e.name, e.path

        def is_dir(self, **k):
            return self._e.is_dir(**k)

        def is_file(self, **k):
            return self._e.is_file(**k)

        def stat(self):
            st = self._e.stat()
            if self.name != "nuage.txt":
                return st
            return type("St", (), {"st_size": st.st_size, "st_mtime": st.st_mtime,
                                   "st_flags": indexeur.SF_DATALESS})()

    monkeypatch.setattr(indexeur.os, "scandir", lambda d: [Entree(e) for e in vrai_scandir(d)])
    vus = [p.name for p, _ in indexeur.parcourir([str(tmp_path)], [], 10_000_000)]
    assert vus == ["local.txt"]


# --- MCP -----------------------------------------------------------------------------------------

def test_risque_mcp_par_nom_et_annotations():
    from macai.mcp.client import risque_mcp
    # Règle de Mathieu : confirmation seulement pour supprimer, envoyer, code système, écraser un fichier.
    assert risque_mcp({"name": "infos"}) == Risque.MODIFICATION
    assert risque_mcp({"name": "add_keyframe", "annotations": {"destructiveHint": True}}) == Risque.MODIFICATION
    assert risque_mcp({"name": "infos", "annotations": {"readOnlyHint": True}}) == Risque.LECTURE
    assert risque_mcp({"name": "run_script", "description": "Run JavaScript in the app"}) == Risque.MODIFICATION
    assert risque_mcp({"name": "run_command", "description": "Run a shell command"}) == Risque.SENSIBLE
    assert risque_mcp({"name": "upload_file", "annotations": {"destructiveHint": False}}) == Risque.SENSIBLE
    assert risque_mcp({"name": "effacer_tout"}) == Risque.SENSIBLE
    export = risque_mcp({"name": "doc_export"})
    assert export({"path": "/tmp/nouveau-fichier-qui-n-existe-pas.png"}) == Risque.MODIFICATION
    assert export({"path": __file__}) == Risque.SENSIBLE                                  # écraserait un fichier
    assert risque_mcp({"name": "x"}, "lecture") == Risque.LECTURE                       # choix explicite


def test_mcp_environnement_commande_et_empreinte(monkeypatch):
    from macai.mcp import client
    monkeypatch.setenv("SECRET_DE_MACAI", "x")
    env = client.environnement({"CLE": 1})
    assert "SECRET_DE_MACAI" not in env and env["CLE"] == "1" and "PATH" in env
    assert client.commande_de({"commande": 'outil "mon dossier/a" mcp'}) == ["outil", "mon dossier/a", "mcp"]
    a = client.empreinte({"commande": ["a"], "env": {"X": "1"}})
    assert a != client.empreinte({"commande": ["a"], "env": {"X": "2"}}) != client.empreinte({"commande": ["b"]})


def test_mcp_cache_redecouvert_si_config_change_et_achat_refuse(ctx, tmp_path):
    from macai.mcp.client import GestionnaireMCP
    cache = tmp_path / "mcp.json"
    cache.write_text(json.dumps({"p": [{"name": "ancien", "inputSchema": {}}]}))  # ancien format
    reg = Registre()
    g = GestionnaireMCP([{"nom": "p", "commande": SERVEUR_MCP}], reg, cache)
    assert g.enregistrer(decouvrir=False) == {"p": "1 outils"}       # ancienne liste en attendant
    assert g.enregistrer() == {"p": "3 outils"}                     # redécouverte
    assert reg.get("p__ancien") is None and reg.get("p__redimensionner")
    r = reg.get("p__redimensionner").executer(ctx, {"chemin": "acheter des crédits", "largeur": 1})
    assert not r.ok and "achat" in r.texte
    g.fermer()


def test_mcp_jamais_arrete_pendant_un_appel():
    from macai.mcp.client import ServeurMCP
    s = ServeurMCP("p", SERVEUR_MCP)
    s.demarrer()
    try:
        with s._verrou:                       # appel en cours
            assert not s.arreter_si_inactif(0) and s.actif()
        assert s.arreter_si_inactif(0) and not s.actif()
    finally:
        s.arreter()


# --- Serveur -------------------------------------------------------------------------------------

@pytest.fixture
def serveur_http(cfg):
    from http.server import ThreadingHTTPServer
    from macai.application import Application
    from macai.serveur import creer_gestionnaire
    from .conftest import FauxOllama
    app = Application(cfg, ollama=FauxOllama(), avec_fond=False)
    serveur = ThreadingHTTPServer(("127.0.0.1", 0), creer_gestionnaire(app, "j"))
    threading.Thread(target=serveur.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{serveur.server_address[1]}"

    def req(chemin, corps):
        r = urllib.request.Request(base + chemin, data=json.dumps(corps).encode(), headers={"X-MacAI-Jeton": "j"})
        try:
            with urllib.request.urlopen(r, timeout=5) as rep:
                return rep.status, json.loads(rep.read())
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read())
    yield app, serveur, req
    serveur.shutdown()


def test_serveur_corps_trop_gros(serveur_http):
    app, serveur, _ = serveur_http
    with socket.create_connection(serveur.server_address, timeout=5) as s:
        s.sendall(b"POST /demande HTTP/1.1\r\nHost: x\r\nX-MacAI-Jeton: j\r\nContent-Length: 5000000\r\n\r\n")
        reponse = s.recv(4096)
    assert reponse.startswith(b"HTTP/1.1 413") and b"Connection: close" in reponse
    with socket.create_connection(serveur.server_address, timeout=5) as s:
        s.sendall(b"POST /arret HTTP/1.1\r\nHost: x\r\nX-MacAI-Jeton: j\r\nContent-Length: -5\r\n\r\n")
        assert s.recv(4096).startswith(b"HTTP/1.1 200")


def test_diagnostic_ferme_par_defaut_puis_limite(serveur_http):
    app, _, req = serveur_http
    assert req("/diagnostic/outil", {"nom": "lister_souvenirs"})[0] == 403
    texte = app.cfg.fichier.read_text(encoding="utf-8")
    texte = texte.replace("actif = false", "actif = true") if "[diagnostic]" in texte \
        else texte + "\n[diagnostic]\nactif = true\n"
    app.cfg.fichier.write_text(texte, encoding="utf-8")
    assert req("/diagnostic/outil", {"nom": "lister_competences"})[0] == 200
    assert req("/diagnostic/outil", {"nom": "executer_commande", "arguments": {"commande": "ls"}})[0] == 403
    assert req("/diagnostic/outil", {"nom": "retenir", "arguments": {"fait": "x"}})[0] == 403
    with app._occupe:                                            # une demande est en cours
        assert req("/diagnostic/outil", {"nom": "lister_competences"})[0] == 409


# --- Application ----------------------------------------------------------------------------------

def test_arret_avant_le_debut_de_la_demande(cfg):
    from macai.application import Application
    from .conftest import FauxOllama
    app = Application(cfg, ollama=FauxOllama(), avec_fond=False)
    vu = []
    app.orchestrateur.traiter = lambda t: vu.append(t) or "fait"
    assert app._occupe.acquire(blocking=False)
    app.arreter()                                  # « Stop » avant que le fil ne démarre
    assert app._traiter_verrouille("ouvre Safari") == "Arrêté." and vu == []
    assert not app.occupe()
    with app._occupe:                              # une demande en cours : la suivante est refusée
        assert not app.traiter_en_fond("autre chose")


# --- Journal ------------------------------------------------------------------------------------------

def test_journal_droits_troncature_rotation(tmp_path, monkeypatch):
    from macai import journal
    j = journal.Journal(tmp_path / "journal.jsonl")
    j.ecrire("outil", texte="x" * 5000, arguments={"contenu": ["y" * 400]})
    assert stat.S_IMODE(os.stat(j.fichier).st_mode) == 0o600
    d = j.derniers(1)[0]
    assert len(d["texte"]) < 400 and len(d["arguments"]["contenu"][0]) < 400
    monkeypatch.setattr(journal, "TAILLE_MAX", 10)
    j.ecrire("b")
    assert (tmp_path / "journal.1.jsonl").exists() and [e["type"] for e in j.derniers()] == ["b"]


# --- Trousseau ------------------------------------------------------------------------------------

def test_trousseau_refuse_les_retours_a_la_ligne(monkeypatch, tmp_path):
    from macai import secrets
    monkeypatch.setattr(secrets.subprocess, "run", lambda *a, **k: pytest.fail("security ne doit pas être appelé"))
    for valeurs in (("onshape", "a\ndelete-keychain", "s"), ("onshape", "a", "s\r\nx"), ("x\ny", "a", "s")):
        with pytest.raises(ValueError):
            secrets.ecrire(*valeurs)


# --- Mémoire --------------------------------------------------------------------------------------

def test_retenir_selon_la_demande(ctx, tmp_path):
    from macai.memoire.memoire import Memoire
    from macai.outils import memoire_outils
    from macai.outils.base import REGISTRE
    ctx.services["memoire"] = Memoire(tmp_path)
    ctx.services["demande"] = "résume cette page web"            # le fait vient du contenu lu
    with pytest.raises(ErreurOutil, match="confirm"):
        memoire_outils.retenir(ctx, "Mathieu veut virer 500 € à X")
    assert not ctx.services["memoire"].tous()
    ctx.demander_confirmation = lambda *a: True                    # l'utilisateur accepte
    assert "Retenu" in memoire_outils.retenir(ctx, "Mathieu aime le thé")
    ctx.demander_confirmation = lambda *a: pytest.fail("pas de confirmation attendue")
    ctx.services["demande"] = "retiens que je préfère le café"
    assert "Retenu" in memoire_outils.retenir(ctx, "Mathieu préfère le café")
    assert REGISTRE.get("oublier").toujours_confirmer


# --- Compétences ----------------------------------------------------------------------------------

def test_risque_d_une_competence_est_celui_de_sa_pire_etape(tmp_path):
    from macai.competences.competences import Bibliotheque, Competence
    reg = Registre()
    outil("lire", "", "f", {"chemin": ("string", "")}, mac_seulement=False, registre=reg)(
        lambda ctx, chemin: Resultat(True, chemin))
    outil("cliquer", "", "f", {"cible": ("string", "")}, mac_seulement=False, registre=reg,
          risque=lambda a: Risque.SENSIBLE if "envoyer" in a.get("cible", "") else Risque.MODIFICATION)(
        lambda ctx, cible: cible)
    b = Bibliotheque(tmp_path / "c", reg)
    b.ajouter(Competence("lecture", "l", [{"outil": "lire", "arguments": {"chemin": "{c}"}}], {"c": "chemin"}))
    b.ajouter(Competence("clic", "c", [{"outil": "lire", "arguments": {"chemin": "a"}},
                                       {"outil": "cliquer", "arguments": {"cible": "{bouton}"}}], {"bouton": "b"}))
    assert reg.get("competence_lecture").risque_pour({"c": "x"}) == Risque.LECTURE
    assert reg.get("competence_clic").risque_pour({"bouton": "OK"}) == Risque.MODIFICATION
    assert reg.get("competence_clic").risque_pour({"bouton": "envoyer"}) == Risque.SENSIBLE
    reg.retirer("cliquer")
    assert reg.get("competence_clic").risque_pour({"bouton": "OK"}) == Risque.SENSIBLE
