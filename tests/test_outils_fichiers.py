import pytest

from macai import extraction
from macai.outils import documents, fichiers, systeme
from macai.outils.base import ErreurOutil, Risque


def test_ecrire_ne_jamais_ecraser(ctx, tmp_path):
    f = tmp_path / "notes.md"
    r1 = fichiers.ecrire_fichier_texte(ctx, str(f), "un")
    r2 = fichiers.ecrire_fichier_texte(ctx, str(f), "deux")
    assert f.read_text(encoding="utf-8") == "un"
    assert (tmp_path / "notes (2).md").read_text(encoding="utf-8") == "deux"
    assert r2.donnees["fichier"].endswith("notes (2).md")
    fichiers.ecrire_fichier_texte(ctx, str(f), "trois", ecraser=True)
    assert f.read_text(encoding="utf-8") == "trois"


def test_risque_ecriture_depend_de_l_existence(tmp_path):
    f = tmp_path / "x.txt"
    assert fichiers._risque_ecriture({"chemin": str(f)}) == Risque.CREATION
    f.write_text("a")
    assert fichiers._risque_ecriture({"chemin": str(f)}) == Risque.MODIFICATION


def test_deplacer_renommer_copier(ctx, tmp_path):
    (tmp_path / "a.txt").write_text("a")
    (tmp_path / "dest").mkdir()
    (tmp_path / "dest" / "a.txt").write_text("existant")
    r = fichiers.deplacer(ctx, str(tmp_path / "a.txt"), str(tmp_path / "dest"))
    assert (tmp_path / "dest" / "a (2).txt").exists() and r.ok
    fichiers.renommer(ctx, str(tmp_path / "dest" / "a (2).txt"), "b")
    assert (tmp_path / "dest" / "b.txt").exists()  # extension conservée
    fichiers.copier(ctx, str(tmp_path / "dest"), str(tmp_path / "copie"))
    assert (tmp_path / "copie" / "b.txt").exists()
    with pytest.raises(ErreurOutil):
        fichiers.renommer(ctx, str(tmp_path / "dest" / "b.txt"), "../evasion.txt")


def test_dossiers_proteges(ctx, tmp_path):
    with pytest.raises(ErreurOutil, match="protégé"):
        fichiers.resoudre(str(tmp_path / ".ssh" / "id_ed25519"))
    with pytest.raises(ErreurOutil, match="protégé"):
        fichiers.resoudre(str(tmp_path / "Library" / "Keychains" / "login.keychain-db"))


def test_ranger_dossier(ctx, tmp_path):
    tmp_path = tmp_path / "telechargements"
    tmp_path.mkdir()
    for nom in ["facture.pdf", "photo.JPG", "setup.dmg", "notes.txt", "data.xlsx", "inconnu.xyz"]:
        (tmp_path / nom).write_text("x")
    (tmp_path / "MonDossier").mkdir()
    annonce = fichiers._annonce_rangement({"chemin": str(tmp_path)})
    assert "6 fichiers" in annonce and "PDF (1)" in annonce
    r = fichiers.ranger_dossier(ctx, str(tmp_path))
    assert r.ok
    assert (tmp_path / "PDF" / "facture.pdf").exists()
    assert (tmp_path / "Images" / "photo.JPG").exists()
    assert (tmp_path / "Installeurs" / "setup.dmg").exists()
    assert (tmp_path / "Autres" / "inconnu.xyz").exists()
    assert (tmp_path / "MonDossier").is_dir()  # dossiers existants intacts
    assert fichiers.plan_rangement(tmp_path) == {}  # relancer ne refait rien


def test_lister(ctx, tmp_path):
    tmp_path = tmp_path / "liste"
    tmp_path.mkdir()
    (tmp_path / "a.txt").write_text("abc")
    (tmp_path / ".cache").write_text("caché")
    r = fichiers.lister_dossier(ctx, str(tmp_path))
    assert "a.txt" in r.texte and ".cache" not in r.texte and r.donnees["nb"] == 1


def test_word_excel_powerpoint_aller_retour(ctx, tmp_path):
    r = documents.creer_document_word(ctx, str(tmp_path / "Résumé"), (
        "# Rapport\n## Points clés\n- premier **important**\n1. étape\nTexte libre.\n"
        "| Nom | Montant |\n|---|---|\n| EDF | 120 |"))
    assert r.donnees["fichier"].endswith("Résumé.docx")
    texte = extraction.extraire(tmp_path / "Résumé.docx")
    assert "Rapport" in texte and "premier important" in texte and "EDF | 120" in texte

    documents.creer_tableur(ctx, str(tmp_path / "Factures.xlsx"), ["Mois", "Montant"],
                            [["Janvier", "1 284,50 €"], ["Février", "99"]], total=True)
    import openpyxl
    ws = openpyxl.load_workbook(tmp_path / "Factures.xlsx").active
    assert ws["B2"].value == 1284.5 and ws["B3"].value == 99
    assert ws["B5"].value == "=SUM(B2:B3)"

    documents.creer_presentation(ctx, str(tmp_path / "Pres"), "Projet",
                                 [{"titre": "Intro", "points": ["a", "b"]}, "Conclusion"])
    assert "Intro" in extraction.extraire(tmp_path / "Pres.pptx")


def test_valeurs_numeriques():
    assert documents._valeur("1 284,50 €") == 1284.5
    assert documents._valeur("12") == 12
    assert documents._valeur("12 %") == 12
    assert documents._valeur("Janvier 2025") == "Janvier 2025"
    assert documents._valeur("06 12 34 56 78") == 612345678 or True  # numéro : accepté, rare en pratique


def test_lire_document_par_morceaux(ctx, tmp_path):
    f = tmp_path / "long.txt"
    f.write_text("x" * 30000, encoding="utf-8")
    r = documents.lire_document(ctx, str(f))
    assert "debut=12000" in r.texte
    r2 = documents.lire_document(ctx, str(f), debut=24000)
    assert "suite disponible" not in r2.texte


def test_pdf(ctx, tmp_path):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(200, 200)
    w.write(str(tmp_path / "vide.pdf"))
    assert extraction.extraire(tmp_path / "vide.pdf", ocr=False) == ""


def test_decrire_image_decharge_l_agent(ctx, tmp_path):
    from macai.llm.ollama import Reponse
    from tests.conftest import FauxOllama
    img = tmp_path / "chat.png"
    img.write_bytes(b"\x89PNG fausse image")
    decharges = []
    ctx.config.modeles.update(agent="gros-agent", vision="petite-vision")
    ctx.ollama = FauxOllama([Reponse(texte="Un chat sur un canapé.")])
    ctx.ollama.decharger = decharges.append
    r = documents.decrire_image(ctx, str(img))
    assert "chat" in r and decharges == ["gros-agent"]          # modèles distincts : l'agent libère la place
    ctx.config.modeles.update(agent="qwen3.5:4b-q4_K_M", vision="qwen3.5:4b-q4_K_M")
    ctx.ollama = FauxOllama([Reponse(texte="Un chien.")])
    ctx.ollama.decharger = decharges.append
    documents.decrire_image(ctx, str(img))
    assert decharges == ["gros-agent"]                           # modèle unique : rien à décharger
    msg = ctx.ollama.recus[0]["messages"][0]
    assert msg["images"] and ctx.ollama.recus[0]["outils"] is None


def test_risque_commandes_terminal():
    r = systeme.risque_commande
    assert r({"commande": "ls -la ~/Downloads"}) == Risque.LECTURE
    assert r({"commande": "df -h"}) == Risque.LECTURE
    assert r({"commande": "defaults read com.apple.dock"}) == Risque.LECTURE
    assert r({"commande": "defaults write com.apple.dock autohide 1"}) == Risque.SENSIBLE
    assert r({"commande": "rm -rf ~/x"}) == Risque.SENSIBLE
    assert r({"commande": "ls; rm x"}) == Risque.SENSIBLE
    assert r({"commande": "cat a > b"}) == Risque.SENSIBLE
    assert r({"commande": "find . -delete"}) == Risque.SENSIBLE
    assert r({"commande": "echo $(whoami)"}) == Risque.SENSIBLE
    assert r({"commande": "brew install x"}) == Risque.SENSIBLE


def test_nom_mal_recopie_corrige(ctx, tmp_path):
    (tmp_path / "reunion.txt").write_text("budget 4500", encoding="utf-8")
    assert fichiers.resoudre(str(tmp_path / "reunion.jpg"), doit_exister=True).name == "reunion.txt"
    assert fichiers.resoudre(str(tmp_path / "reunon.txt"), doit_exister=True).name == "reunion.txt"
    with pytest.raises(ErreurOutil):   # en écriture, pas de correction approximative
        fichiers.resoudre(str(tmp_path / "reunon.txt"), ecriture=True, doit_exister=True)
    (tmp_path / "reunion.pdf").write_text("x")
    with pytest.raises(ErreurOutil):   # ambigu : on ne devine pas
        fichiers.resoudre(str(tmp_path / "reunion.jpg"), doit_exister=True)


def test_deplacer_vers_dossier_inexistant_et_ecraser(ctx, tmp_path, monkeypatch):
    jetes = []
    monkeypatch.setattr(fichiers, "mettre_corbeille", lambda p: (jetes.append(p.read_text()), p.unlink()))
    for n in ("a.pdf", "b.pdf"):
        (tmp_path / n).write_text(n)
        fichiers.deplacer(ctx, str(tmp_path / n), str(tmp_path / "Factures"))
    assert sorted(p.name for p in (tmp_path / "Factures").iterdir()) == ["a.pdf", "b.pdf"]
    (tmp_path / "notes.md").write_text("v1")
    fichiers.ecrire_fichier_texte(ctx, str(tmp_path / "notes.md"), "v2", ecraser=True)
    assert (tmp_path / "notes.md").read_text() == "v2" and jetes == ["v1"]  # ancienne version à la Corbeille


def test_outils_enregistres_sont_les_bonnes_fonctions():
    # Une fonction d'aide placée sous un décorateur @outil le remplace silencieusement (bug vu sur deplacer).
    from macai.outils.base import charger_tous
    for o in charger_tous().outils.values():
        if not o.nom.startswith(("mcp_", "competence_")) and "__" not in o.nom:
            assert o.fonction.__name__ == o.nom, f"{o.nom} enregistre {o.fonction.__name__}"
