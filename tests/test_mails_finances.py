"""Mails et notes dans l'index, tri des mails, dépenses, classement des papiers — tout simulé (rien sur le vrai Mac)."""
import os
import time
from datetime import date, datetime

import pytest

from macai import annulation, classement
from macai.mac import commun
from macai.outils import finances as fin
from macai.outils import organisation as org
from macai.outils import recherche_outils
from macai.outils.base import REGISTRE, ErreurOutil, Risque, charger_tous
from macai.recherche import indexeur as ix
from macai.recherche.index import Index

from .conftest import FauxOllama

US, RS = "\x1f", "\x1e"


# --- Montants, dates, marchands -------------------------------------------------------------------------------

@pytest.mark.parametrize("texte, attendu", [
    ("1 284,50 €", 1284.5), ("84.20 EUR", 84.2), ("€ 12,00", 12.0), ("9,99 euros", 9.99), ("12 €", 12.0),
    ("1 284,50 €", 1284.5), ("1.284,50 €", 1284.5), ("1,284.50 EUR", 1284.5), ("pas de prix", None),
])
def test_lire_montant(texte, attendu):
    assert fin.lire_montant(texte) == attendu


def test_montant_total_prefere_le_libelle_au_plus_grand():
    facture = ("EDF Facture n° 123\nTotal HT 70,17 €\nTVA 20 % 14,03 €\nTotal TTC 84,20 €\n"
               "Consommation annuelle estimée : 1 450,00 €")
    assert fin.montant_total(facture) == 84.2
    assert fin.montant_total("Net à payer : 1 284,50") == 1284.5        # devise en tête de colonne
    assert fin.montant_total("Article 3,20 €\nArticle 12,00 €") == 12.0  # sans libellé : le plus grand
    assert fin.euros(1284.5) == "1 284,50 €"


@pytest.mark.parametrize("texte, attendu", [
    ("le 15/10/2026", date(2026, 10, 15)), ("15 octobre 2026", date(2026, 10, 15)),
    ("1er oct. 2026", date(2026, 10, 1)), ("2026-10-15", date(2026, 10, 15)), ("15-10-26", date(2026, 10, 15)),
    ("15 Février 2026", date(2026, 2, 15)), ("31/02/2026", None),
])
def test_lire_date(texte, attendu):
    assert fin.lire_date(texte) == attendu


def test_date_document_suit_le_libelle():
    t = "Période du 01/09/2026 au 30/09/2026\nDate de facture : 05/10/2026"
    assert fin.date_document(t) == date(2026, 10, 5)


def test_marchands_et_categories():
    assert fin.trouver_marchand("TICKET CARREFOUR MARKET\npayé par Apple Pay") == ("Carrefour", "courses")
    assert fin.trouver_marchand("Total TTC 12,00 €") is None              # « Total » n'est pas TotalEnergies
    m = fin.trouver_marchand("TotalEnergies Relais\nSP98 42,10 L")
    assert fin.categorie_depense("SP98 42,10 litres", m) == "carburant"
    assert fin.categorie_depense("Abonnement Freebox", fin.trouver_marchand("Freebox")) == "télécom"
    assert fin.normaliser_categorie("essence") == "carburant"
    assert fin.normaliser_categorie("Électricité") == "énergie"
    assert fin.normaliser_categorie("") == ""
    with pytest.raises(ErreurOutil):
        fin.normaliser_categorie("licornes")


@pytest.mark.parametrize("periode, debut, fin_, libelle", [
    ("ce mois", date(2026, 10, 1), date(2026, 11, 1), "octobre 2026"),
    ("ce mois-ci", date(2026, 10, 1), date(2026, 11, 1), "octobre 2026"),
    ("le mois dernier", date(2026, 9, 1), date(2026, 10, 1), "septembre 2026"),
    ("novembre", date(2025, 11, 1), date(2025, 12, 1), "novembre 2025"),
    ("mars 2026", date(2026, 3, 1), date(2026, 4, 1), "mars 2026"),
    ("2026-02", date(2026, 2, 1), date(2026, 3, 1), "février 2026"),
    ("cette année", date(2026, 1, 1), date(2027, 1, 1), "2026"),
    ("l'année dernière", date(2025, 1, 1), date(2026, 1, 1), "2025"),
])
def test_bornes_periode(periode, debut, fin_, libelle):
    assert fin.bornes_periode(periode, date(2026, 10, 16)) == (debut, fin_, libelle)


def test_bornes_periode_glissante_et_refus():
    d, f, _ = fin.bornes_periode("3 derniers mois", date(2026, 10, 16))
    assert d == date(2026, 7, 18) and f == date(2026, 10, 17)
    with pytest.raises(ErreurOutil, match="Période non comprise"):
        fin.bornes_periode("quand il pleut")


# --- Dépenses depuis l'index -----------------------------------------------------------------------------------

def _index_depenses(tmp_path):
    index = Index(tmp_path / "index.db", None)
    docs = tmp_path / "docs"
    docs.mkdir()
    maintenant = time.time()
    fichiers = {
        "facture_edf.pdf": "EDF\nFacture d'électricité\nDate de facture : 05/10/2026\nTotal TTC 84,20 €",
        "ticket_total.pdf": "TotalEnergies Relais Toulon\nGazole 38,5 L\nTOTAL 62,30 EUR\nCB le 12/10/2026",
        "ticket_esso.pdf": "ESSO EXPRESS\nSP95-E10\nMontant 45,00 €\nTicket du 14/10/2026",
        "ticket_lidl.pdf": "LIDL\nTicket de caisse\nTotal à payer 1 284,50 €\n15/10/2026",
        "devis_cuisine.pdf": "Devis cuisine\nTotal TTC 8 000,00 €\n10/10/2026",
        "ancien.pdf": "Carrefour ticket Total 10,00 € 15/09/2026",
    }
    for nom, texte in fichiers.items():
        p = docs / nom
        p.write_text(texte, encoding="utf-8")
        index.enregistrer(p, maintenant, len(texte), "pdf", texte)
    # Le mail qui annonce la même facture EDF : compté une seule fois.
    index.enregistrer_entree("mail://abc@edf.fr", datetime(2026, 10, 6, 8).timestamp(), 0, "mail",
                             "Mail « Votre facture est disponible » — de EDF <noreply@edf.fr> — reçu le 06/10/2026",
                             "Votre facture d'un montant de 84,20 € est disponible.")
    return index


def test_releve_depenses(ctx, tmp_path, monkeypatch):
    ctx.services["index"] = _index_depenses(tmp_path)
    monkeypatch.setattr(fin, "bornes_periode", lambda p, a=None: (date(2026, 10, 1), date(2026, 11, 1),
                                                                  "octobre 2026"))
    r = fin.releve_depenses(ctx, "ce mois")
    assert r.ok and r.donnees["nombre"] == 4, r.texte          # devis exclu, mail EDF dédoublonné, septembre exclu
    assert r.donnees["total"] == pytest.approx(84.2 + 62.3 + 45 + 1284.5)
    assert "carburant : 107,30 € (2)" in r.texte and "facture_edf.pdf" in r.texte
    essence = fin.releve_depenses(ctx, "ce mois", "essence")
    assert essence.donnees["total"] == pytest.approx(107.3) and "Lidl" not in essence.texte


def test_exporter_depenses(ctx, tmp_path, monkeypatch):
    ctx.services["index"] = _index_depenses(tmp_path)
    monkeypatch.setattr(fin, "bornes_periode", lambda p, a=None: (date(2026, 10, 1), date(2026, 11, 1),
                                                                  "octobre 2026"))
    from macai.outils import documents
    monkeypatch.setattr(documents, "destination", lambda c, e: tmp_path / ("sortie" + e))
    r = fin.exporter_depenses(ctx, "ce mois")
    assert r.ok and (tmp_path / "sortie.xlsx").exists() and "1 476,00 €" in r.texte
    import openpyxl
    ws = openpyxl.load_workbook(tmp_path / "sortie.xlsx").active
    assert ws.cell(1, 4).value == "Montant (€)" and ws.cell(2, 4).value == 84.2


# --- Mails et notes dans l'index -------------------------------------------------------------------------------

class FauxMail:
    """Répond aux scripts de SourcesApple comme Mail et Notes le feraient."""

    def __init__(self, mails, notes):
        self.mails, self.notes, self.scripts = mails, notes, []

    def __call__(self, script, delai=30):
        self.scripts.append(script)
        if "message id of messages 1 thru" in script:
            return RS.join(f"{m['boite']}{US}{i}{US}{m['id']}{US}{m['age']}" for i, m in enumerate(self.mails, 1))
        if "content of m" in script:
            return RS.join(f"{m['boite']}{US}{m['id']}{US}{m['de']}{US}{m['sujet']}{US}{m['age']}{US}{US}{m['texte']}"
                           for m in self.mails if f'"{m["id"]}"' in script)
        if "id of every note" in script:
            return RS.join(f"{n['id']}{US}{n['age']}{US}{n['nom']}" for n in self.notes)
        if "plaintext of n" in script:
            return RS.join(f"{n['id']}{US}{n['nom']}\n{n['texte']}" for n in self.notes if f'"{n["id"]}"' in script)
        return ""


def test_mails_et_notes_indexes_et_retrouves(ctx, tmp_path):
    index = Index(tmp_path / "index.db", FauxOllama())
    faux = FauxMail(
        [{"boite": "recus", "id": "m1@exemple.fr", "de": "Paul Martin <paul@gmail.com>", "sujet": "Devis cuisine",
          "age": 3600, "texte": "Voici le devis pour la cuisine en chêne."},
         {"boite": "envoyes", "id": "m2@exemple.fr", "de": "moi", "sujet": "Vacances", "age": 7200,
          "texte": "On part à Annecy en juillet."}],
        [{"id": "x-coredata://N1", "age": 100, "nom": "Courses", "texte": "lait, farine, oeufs"}])
    sources = ix.SourcesApple(osascript=faux)
    assert sources.indexer(index, sortie=lambda *a: None) == {"mails": (2, 0), "notes": (1, 0)}
    assert set(index.chemins("mail://")) == {"mail://m1@exemple.fr", "mail://m2@exemple.fr"}
    # Second passage : rien de nouveau, aucun détail relu.
    faux.scripts.clear()
    assert sources.indexer(index, sortie=lambda *a: None) == {"mails": (0, 0), "notes": (0, 0)}
    assert not any("content of m" in s or "plaintext" in s for s in faux.scripts)
    # Note modifiée, mail disparu.
    faux.notes[0].update(age=1, texte="lait, chocolat")
    del faux.mails[1]
    sources.indexer(index, sortie=lambda *a: None, maintenant=time.time() + 60)
    assert list(index.chemins("mail://")) == ["mail://m1@exemple.fr"]
    ctx.services["index"] = index
    r = recherche_outils.chercher_partout(ctx, "devis cuisine")
    assert "mail://m1@exemple.fr" in r.texte and "Paul Martin" in r.texte and "Devis cuisine" in r.texte
    assert "chocolat" in recherche_outils.chercher_partout(ctx, "chocolat", sources="notes").texte
    assert "mail://" not in recherche_outils.rechercher_fichiers(ctx, "devis cuisine").texte
    assert "Mails : 1, notes : 1" in recherche_outils.etat_index(ctx)


def test_indexeur_ne_prend_pas_les_mails_pour_des_fichiers_disparus(cfg, tmp_path):
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "a.txt").write_text("bonjour", encoding="utf-8")
    cfg.recherche["dossiers"] = [str(docs)]
    index = Index(tmp_path / "index.db", None)
    index.enregistrer_entree("mail://z@y", time.time(), 0, "mail", "Mail « z »", "texte")
    i = ix.Indexeur(cfg, index, sortie=lambda *a: None)
    assert i.sources is None                       # dossier de test : ni Mail ni Notes
    assert i.executer(force=True)["retires"] == 0 and index.connu("mail://z@y")


def test_mail_ferme_rien_retire(tmp_path):
    index = Index(tmp_path / "index.db", None)
    index.enregistrer_entree("mail://z@y", time.time(), 0, "mail", "Mail « z »", "texte")
    sources = ix.SourcesApple(osascript=lambda s, delai=30: "FERME")
    assert sources.indexer(index, sortie=lambda *a: None) == {"mails": (0, 0), "notes": (0, 0)}
    assert index.connu("mail://z@y")
    assert 'is not running then return "FERME"' in ix.script_lister_mails()
    assert "is not running" not in ix.script_lister_mails(lancer=True)


def test_url_message():
    assert recherche_outils.url_message("mail://abc.123@mail.gmail.com") == "message://%3Cabc.123@mail.gmail.com%3E"
    assert recherche_outils.genres_de_sources("mails") == (["mail"], None)
    assert recherche_outils.genres_de_sources("fichiers") == (None, ["mail", "note"])
    assert recherche_outils.genres_de_sources("tout") == (None, None)


# --- Tri des mails ---------------------------------------------------------------------------------------------

@pytest.fixture
def scripts(monkeypatch):
    journal = {"as": [], "reponses": []}

    def faux_osascript(script, javascript=False, delai=30):
        journal["as"].append(script)
        return journal["reponses"].pop(0) if journal["reponses"] else ""

    monkeypatch.setattr(commun, "osascript", faux_osascript)
    monkeypatch.setitem(org._CONTACTS, "quand", time.time())
    monkeypatch.setitem(org._CONTACTS, "adresses", {"marie@exemple.fr"})
    return journal


def _brut(*mails):
    return RS.join(US.join(map(str, m)) for m in mails)


def test_classer_un_mail_regles():
    contacts = {"marie@exemple.fr"}
    c = lambda **x: org.classer_un_mail({"de": "", "sujet": "", "extrait": "", "desabonnement": False, **x}, contacts)[0]
    assert c(de="Marie <marie@exemple.fr>", sujet="Dîner", extrait="Tu es dispo samedi ?") == "a_traiter"
    assert c(de="Marie <marie@exemple.fr>", sujet="Photos", extrait="Les photos de la fête.") == "a_lire"
    assert c(de="Zalando <news@zalando.fr>", sujet="-40 % sur tout", desabonnement=True) == "pubs"
    assert c(de="Fnac <offres@fnac.com>", sujet="Découvrez nos bons plans") == "pubs"
    assert c(de="Amazon <confirmation@amazon.fr>", sujet="Votre commande a été expédiée") == "notifications"
    assert c(de="Google <no-reply@accounts.google.com>", sujet="Alerte de sécurité") == "notifications"
    assert c(de="Jean <jean.dupont@orange.fr>", sujet="Re: devis", extrait="Peux-tu confirmer ?") == "a_traiter"
    assert c(de="Inconnu <x@societe.fr>", sujet="Information") == "a_lire"


def test_trier_mails_groupe_et_numerote(ctx, scripts):
    scripts["reponses"] = [_brut(
        ("id1@x", "Zalando <news@zalando.fr>", "Soldes -40 %", 600, "1", "Profitez"),
        ("id2@x", "Marie <marie@exemple.fr>", "Dîner samedi", 60, "0", "Tu es dispo samedi soir ?"),
        ("id3@x", "Amazon <ship@amazon.fr>", "Votre colis est expédié", 120, "0", "Livraison demain"))]
    r = org.trier_mails(ctx, 30)
    assert "whose read status is false" in scripts["as"][0] and "list-unsubscribe" in scripts["as"][0]
    t = r.texte
    assert t.index("À traiter (1)") < t.index("Notifications (1)") < t.index("Pubs et newsletters (1, désinscription")
    assert "1. Marie — " in t and "3. Zalando — " in t
    assert r.donnees["categories"] == {"a_traiter": 1, "a_lire": 0, "notifications": 1, "pubs": 1}
    assert org.DERNIER_TRI[3]["id"] == "id1@x"


def test_classer_mail(ctx, scripts):
    org.DERNIER_TRI.clear()
    org.DERNIER_TRI[2] = {"id": "id1@x", "sujet": "Soldes"}
    scripts["reponses"] = [f"INBOX{US}1{US}Soldes"]
    r = org.classer_mail(ctx, "2", "Newsletters")
    assert 'message id is "id1@x"' in scripts["as"][0] and 'mailbox "Newsletters"' in scripts["as"][0]
    assert "rangé dans « Newsletters » (boîte créée)" in r and "INBOX" in r
    with pytest.raises(ErreurOutil, match="supprimer"):
        org.classer_mail(ctx, "2", "Corbeille")
    with pytest.raises(ErreurOutil, match="trier_mails"):
        org.classer_mail(ctx, "9", "Archives")
    assert org.identifiant_mail("mail://<abc@d>")[0] == "abc@d"


def test_risques_des_nouveaux_outils():
    charger_tous()
    assert REGISTRE.get("trier_mails").risque_pour({}) == Risque.LECTURE
    assert REGISTRE.get("classer_mail").risque_pour({}) == Risque.MODIFICATION
    assert REGISTRE.get("releve_depenses").risque_pour({}) == Risque.LECTURE
    assert REGISTRE.get("exporter_depenses").risque_pour({}) == Risque.CREATION
    assert REGISTRE.get("classer_document").risque_pour({}) == Risque.MODIFICATION
    assert REGISTRE.get("classer_telechargements").risque_pour({}) == Risque.MODIFICATION
    assert REGISTRE.get("chercher_partout").risque_pour({}) == Risque.LECTURE


# --- Classement des papiers ------------------------------------------------------------------------------------

FACTURE = ("EDF\nFacture d'électricité n° de facture 1234\nDate de facture : 05/10/2026\nTotal HT 70,17 €\n"
           "TVA 14,03 €\nTotal TTC 84,20 €")
RELEVE = ("Crédit Agricole\nRelevé de compte n° 10\nAncien solde au 30/09/2026 1 200,00 €\n"
          "Nouveau solde au 31/10/2026 1 350,00 €")
PAIE = "Bulletin de paie\nOctobre 2026\nSalaire de base 2 500,00\nNet imposable 1 980,00\nEmployeur : ACME"


def _fausse_extraction(monkeypatch, textes):
    monkeypatch.setattr(classement.extraction, "extraire", lambda p, limite=0, ocr=True: textes.get(p.name, ""))


def test_proposer_classement(tmp_path, monkeypatch):
    racine = tmp_path / "Administratif"
    for nom in ("scan_001.pdf", "releve.pdf", "paie.pdf", "vacances.pdf", "photo.zip"):
        (tmp_path / nom).write_bytes(b"x")
    _fausse_extraction(monkeypatch, {"scan_001.pdf": FACTURE, "releve.pdf": RELEVE, "paie.pdf": PAIE,
                                     "vacances.pdf": "Nous sommes allés à la plage, il faisait beau tout le temps."})
    p = classement.proposer_classement(tmp_path / "scan_001.pdf", racine=racine)
    assert p["type"] == "facture" and p["nom"] == "2026-10 EDF facture 84,20 €.pdf"
    assert p["destination"] == str(racine / "Factures" / "2026" / p["nom"]) and p["confiance"] >= classement.SEUIL
    r = classement.proposer_classement(tmp_path / "releve.pdf", racine=racine)
    assert r["type"] == "releve" and r["nom"] == "2026-09 Crédit Agricole relevé bancaire.pdf"
    assert "Banque" in r["destination"]
    assert classement.proposer_classement(tmp_path / "paie.pdf", racine=racine)["dossier_type"] == "Fiches de paie"
    assert classement.proposer_classement(tmp_path / "vacances.pdf", racine=racine) is None
    assert classement.proposer_classement(tmp_path / "photo.zip", racine=racine) is None


def test_classer_sans_ecraser_et_annulable(ctx, tmp_path, monkeypatch):
    racine = tmp_path / "Administratif"
    src = tmp_path / "scan_001.pdf"
    src.write_bytes(b"pdf")
    _fausse_extraction(monkeypatch, {"scan_001.pdf": FACTURE})
    occupe = racine / "Factures" / "2026" / "2026-10 EDF facture 84,20 €.pdf"
    occupe.parent.mkdir(parents=True)
    occupe.write_bytes(b"autre")
    texte = classement.classer(ctx, src, racine=racine)
    arrivee = occupe.with_name("2026-10 EDF facture 84,20 € (2).pdf")
    assert "Classé" in texte and arrivee.read_bytes() == b"pdf" and occupe.read_bytes() == b"autre"
    assert not src.exists()
    r = annulation.annuler(ctx.config.chemin("annulations.json"))
    assert r["ok"] and src.read_bytes() == b"pdf" and not arrivee.exists()


def test_verifier_telechargements(ctx, tmp_path, monkeypatch):
    dl, racine = tmp_path / "Downloads", tmp_path / "Administratif"
    dl.mkdir()
    maintenant = time.time()
    fichiers = {"facture.pdf": (FACTURE, 600), "trop_recent.pdf": (FACTURE, 30), "vieux.pdf": (FACTURE, 3 * 86400),
                "flou.pdf": ("Merci pour votre visite, à bientôt. Signature du client.", 600),
                "encours.pdf.crdownload": (FACTURE, 600)}
    for nom, (_, age) in fichiers.items():
        (dl / nom).write_bytes(b"x")
        os.utime(dl / nom, (maintenant - age, maintenant - age))
    _fausse_extraction(monkeypatch, {n: t for n, (t, _) in fichiers.items()})
    vus: set = set()
    faits = classement.verifier_telechargements(ctx, vus, dossier=dl, racine=racine, maintenant=maintenant)
    assert len(faits) == 1 and "facture.pdf" in faits[0]
    assert (racine / "Factures" / "2026").is_dir() and (dl / "trop_recent.pdf").exists() and (dl / "vieux.pdf").exists()
    assert (dl / "flou.pdf").exists()
    # Déjà vus : rien n'est réexaminé.
    assert classement.verifier_telechargements(ctx, vus, dossier=dl, racine=racine, maintenant=maintenant) == []


def test_classer_dossier_bilan(ctx, tmp_path, monkeypatch):
    dl, racine = tmp_path / "Downloads", tmp_path / "Administratif"
    dl.mkdir()
    for nom in ("a.pdf", "b.pdf", "c.pdf"):
        (dl / nom).write_bytes(b"x")
    _fausse_extraction(monkeypatch, {"a.pdf": FACTURE, "b.pdf": "Contrat de location. Signature.",
                                     "c.pdf": "Liste de films à voir ce week-end avec les enfants."})
    bilan = classement.classer_dossier(ctx, dl, racine=racine)
    assert len(bilan["classes"]) == 1 and bilan["inconnus"] == ["c.pdf"]
    texte = classement.resume_bilan(bilan, dl)
    assert "1 document(s) classé(s)" in texte and ("b.pdf" in texte or "Non reconnus" in texte)
